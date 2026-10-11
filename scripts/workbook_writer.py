"""Write the exact-column product workbook with openpyxl; no proprietary runtime."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
SHEET_NAME = '\u5546\u54c1\u6570\u636e'  # product data sheet

def write_workbook(spec):
    headers = list(spec['headers'])
    rows = [list(r) for r in spec['rows']]
    out = spec['output']
    wb = Workbook()
    s = wb.active
    s.title = SHEET_NAME
    s.append(headers)
    for r in rows: s.append(r)
    font = Font(name='Calibri', size=11)
    for row in s.iter_rows():
        for c in row: c.font = font
    for i in range(2, len(rows) + 2):
        s.cell(i, 8).number_format = '0.00'
        s.cell(i, 9).number_format = '0'
    wb.save(out)
    check = load_workbook(out, read_only=True, data_only=False)
    observed = [list(x) for x in check.active.values]
    check.close()
    if observed != [headers] + rows: raise ValueError('WORKBOOK_READBACK_MISMATCH')
    return out
