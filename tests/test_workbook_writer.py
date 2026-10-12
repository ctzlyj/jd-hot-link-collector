import sys, tempfile, unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from openpyxl import load_workbook
from workbook_writer import write_workbook
class WorkbookTests(unittest.TestCase):
    def test_exact_values_font_number_formats_and_unicode(self):
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'p.xlsx'
            headers=['\u5546\u54c1','B','C','D','E','F','G','H','I']
            rows=[['t',None,None,None,None,None,None,12.5,999],[None,None,None,None,None,None,None,7.25,999]]
            write_workbook({'headers':headers,'rows':rows,'output':str(out)})
            wb=load_workbook(out); s=wb.active
            self.assertEqual([[c.value for c in r] for r in s.iter_rows()],[headers]+rows)
            self.assertEqual(s['H2'].number_format,'0.00');self.assertEqual(s['I2'].number_format,'0')
            self.assertEqual(s['A1'].font.name,'Calibri');self.assertEqual(s['A1'].font.size,11)
            wb.close()
    def test_save_error_raises(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(Exception):write_workbook({'headers':['A'],'rows':[[1]],'output':str(Path(d)/'missing-dir'/'x.xlsx')})
    def test_real_product_matrix_with_blank_specs_roundtrips(self):
        from package_contract import HEADERS,table_matrix
        main={'title':'title','attrs':[],'crumbs':[]}
        rows=table_matrix(main,[{'sku':'10000000000001','spec1':'A','spec2':''}],{'10000000000001':{'pp':'8.25'}})
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'product.xlsx'
            write_workbook({'headers':HEADERS,'rows':rows,'output':str(out)})
            wb=load_workbook(out);self.assertEqual(wb.active['H2'].value,8.25);self.assertIsNone(wb.active['G2'].value);wb.close()
if __name__=='__main__':unittest.main()
