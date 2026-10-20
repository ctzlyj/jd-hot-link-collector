"""OpenPyXL workbook author as a local replacement when the shared office runtime is unavailable.
Writes the exact same artifact_workbook_spec.json contract and performs an independent readback."""
import json
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def _norm(v):
    return None if v == '' else v


def write_artifact(spec):
    stage = Path(os.environ['JD_COLLECTOR_STATE_DIR']).resolve()
    output = Path(spec['output']).resolve()
    if not output.is_relative_to(stage / 'packages'):
        raise ValueError('ARTIFACT_OUTPUT_OUTSIDE_TASK_PACKAGE')
    manifest = output.parent / 'artifact_workbook_spec.json'
    manifest.write_text(json.dumps([spec], ensure_ascii=False), encoding='utf-8')
    wb = Workbook()
    ws = wb.active
    ws.title = '商品数据'
    values = [spec['headers']] + spec['rows']
    for r, row in enumerate(values, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c)
            cell.value = _norm(val)
            cell.font = Font(name='Calibri', size=11)
    last = len(values)
    for r in range(2, last + 1):
        ws.cell(row=r, column=8).number_format = '0.00'
        ws.cell(row=r, column=9).number_format = '0'
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    # Independent readback before declaring the artifact written.
    chk = load_workbook(str(output), read_only=True, data_only=False)
    observed = list(chk.active.values)
    chk.close()
    if len(observed) != len(values):
        raise ValueError('AUTHOR_READBACK_MISMATCH')
    for a, e in zip(observed, values):
        if [_norm(x) for x in a] != [_norm(x) for x in e]:
            raise ValueError('AUTHOR_READBACK_MISMATCH')
    return str(output)
